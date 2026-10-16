"""Parse Huawei USG firewall configs to Excel and standard rule JSON.

Refactored from FWConfigParse/hw_ac_parse.py: pandas removed, the workbook
is written directly with openpyxl, and settings come from config.yaml.
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .extractors import (
    extract_address_sets,
    extract_security_policies,
    extract_service_sets,
)

DEFAULT_CONFIG = {
    "input": {"directory": "Config", "pattern": "*", "encoding": "utf-8"},
    "output": {"directory": "Output", "formats": ["excel", "json"]},
}

ADDRESS_SET_COLUMNS = [("set_name", "地址集名称"), ("set_type", "类型"),
                       ("addresses", "地址")]
SERVICE_SET_COLUMNS = [("service_set_name", "服务集名称"), ("services", "服务")]
POLICY_COLUMNS = [("rule_name", "规则名称"), ("source_zone", "源安全区域"),
                  ("destination_zone", "目的安全区域"),
                  ("source_address", "源地址"), ("destination_address", "目的地址"),
                  ("service", "服务"), ("action", "动作")]


def load_config(path: str | None = None) -> dict:
    config_path = Path(path) if path else Path(__file__).with_name("config.yaml")
    if not config_path.exists():
        return DEFAULT_CONFIG
    with open(config_path, encoding="utf-8") as fh:
        loaded = yaml.safe_load(fh) or {}
    return {section: {**DEFAULT_CONFIG[section], **loaded.get(section, {})}
            for section in DEFAULT_CONFIG}


def apply_excel_formatting(worksheet) -> None:
    """Apply borders, bold headers, wrap text and column widths."""
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True, underline="single")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if not cell.value:
                continue
            value = str(cell.value)
            if "\n" in value:
                max_length = max(max_length,
                                 max(len(line) for line in value.split("\n")))
            else:
                max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _write_sheet(workbook, title: str, rows: list, columns: list) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append([header for _, header in columns])
    for row in rows:
        worksheet.append([row.get(key, "") for key, _ in columns])
    apply_excel_formatting(worksheet)


def correlate_policies_with_sets(policies: list, address_sets: list,
                                 service_sets: list) -> list:
    """Replace address-set/service-set references in policies with contents."""
    address_set_map = {s["set_name"]: s["addresses"] for s in address_sets}
    service_set_map = {s["service_set_name"]: s["services"] for s in service_sets}

    correlated_policies = []
    for policy in policies:
        correlated_policy = policy.copy()

        source_addresses = []
        for addr in policy["source_address"].split(", "):
            if addr in address_set_map and address_set_map[addr]:
                source_addresses.extend(address_set_map[addr].split("\n"))
            else:
                source_addresses.append(addr)
        correlated_policy["source_address"] = "\n".join(source_addresses)

        destination_addresses = []
        for addr in policy["destination_address"].split(", "):
            if addr in address_set_map and address_set_map[addr]:
                destination_addresses.extend(address_set_map[addr].split("\n"))
            else:
                destination_addresses.append(addr)
        correlated_policy["destination_address"] = "\n".join(destination_addresses)

        services = []
        for svc in policy["service"].split(", "):
            if svc in service_set_map and service_set_map[svc]:
                services.extend(service_set_map[svc].split("\n"))
            else:
                services.append(svc)
        correlated_policy["service"] = "\n".join(services)

        correlated_policies.append(correlated_policy)

    return correlated_policies


def convert_text(content: str) -> dict:
    """Parse config text; return {sheet_name: (columns, rows)}."""
    lines = content.splitlines()
    address_sets = extract_address_sets(content)
    service_sets = extract_service_sets(lines)
    policies = extract_security_policies(lines)

    sheets = {
        "地址集": (ADDRESS_SET_COLUMNS, address_sets),
        "服务集": (SERVICE_SET_COLUMNS, service_sets),
        "安全策略-原始": (POLICY_COLUMNS, policies),
    }
    if policies and (address_sets or service_sets):
        sheets["安全策略-关联"] = (
            POLICY_COLUMNS,
            correlate_policies_with_sets(policies, address_sets, service_sets),
        )
    return sheets


def process_config_file(file_path: Path, output_dir: Path,
                        encoding: str = "utf-8",
                        formats: list | None = None) -> Path:
    """Parse one Huawei config file and write Excel/JSON outputs."""
    formats = formats or ["excel"]
    content = file_path.read_text(encoding=encoding)
    sheets = convert_text(content)
    output_dir.mkdir(parents=True, exist_ok=True)

    excel_file = output_dir / f"{file_path.stem}.xlsx"
    if "excel" in formats:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, (columns, rows) in sheets.items():
            _write_sheet(workbook, title, rows, columns)
        workbook.save(excel_file)
        print(f"Writing file {excel_file}")
    if "json" in formats:
        json_file = output_dir / f"{file_path.stem}.json"
        export_to_json(json_file, file_path.stem, sheets)
        print(f"Writing file {json_file}")
    return excel_file


def _service_to_protocol_port(service: str) -> tuple:
    """Map one Huawei service string to (protocol, dst_port).

    Handles extractor forms (``tcp 80``, ``tcp 8000 to 8010``) and raw policy
    lines (``protocol tcp destination-port 3389 [to 3390]``).
    """
    raw = re.search(r"protocol\s+(\S+)\s+destination-port\s+(\d+)(?:\s+to\s+(\d+))?",
                    service)
    if raw:
        port = f"{raw.group(2)}-{raw.group(3)}" if raw.group(3) else raw.group(2)
        return raw.group(1), port
    expanded = re.match(r"(\S+)\s+(\d+)(?:\s+to\s+(\d+))?$", service.strip())
    if expanded:
        port = (f"{expanded.group(2)}-{expanded.group(3)}"
                if expanded.group(3) else expanded.group(2))
        return expanded.group(1), port
    if service.strip().lower() in ("icmp", "icmpv6"):
        return service.strip().lower(), ""
    return "", ""


def _policy_to_rule(policy: dict) -> dict:
    """Convert one (correlated) policy row to a standard rule entry."""
    protocols, ports = [], []
    for service in filter(None, policy.get("service", "").split("\n")):
        protocol, port = _service_to_protocol_port(service)
        if protocol and protocol not in protocols:
            protocols.append(protocol)
        if port:
            ports.append(port)
    return {
        "name": policy.get("rule_name", ""),
        "action": policy.get("action", ""),
        "protocol": protocols[0] if protocols else "",
        "src_address": "\n".join(filter(
            None, policy.get("source_address", "").split(", "))),
        "dst_address": "\n".join(filter(
            None, policy.get("destination_address", "").split(", "))),
        "dst_port": "\n".join(ports),
        "source_zone": policy.get("source_zone", ""),
        "destination_zone": policy.get("destination_zone", ""),
        "service": policy.get("service", ""),
    }


def export_to_json(output_path: Path, device_name: str, sheets: dict) -> Path:
    """Export parse results in the standard rule JSON format
    (consumed by Emulation/traffic_emulator; see Emulation/README.md).

    Rules are derived from the correlated policies when available so
    address-set/service-set references are already expanded.
    """
    _, address_sets = sheets.get("地址集", ([], []))
    _, service_sets = sheets.get("服务集", ([], []))
    policy_sheet = sheets.get("安全策略-关联") or sheets.get("安全策略-原始") \
        or ([], [])
    _, policies = policy_sheet

    config_data = {
        "metadata": {
            "device_name": device_name,
            "vendor": "huawei",
            "export_time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        },
        "address_sets": address_sets,
        "service_sets": service_sets,
        "rules": [_policy_to_rule(policy) for policy in policies],
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(config_data, fh, ensure_ascii=False, indent=2)
    return output_path


def run(config: dict) -> None:
    input_dir = Path(config["input"]["directory"])
    pattern = config["input"].get("pattern", "*")
    encoding = config["input"].get("encoding", "utf-8")
    output_dir = Path(config["output"]["directory"])
    formats = config["output"].get("formats", ["excel"])

    if not input_dir.exists():
        print(f"Input directory not found: {input_dir.resolve()}")
        return
    config_files = sorted(p for p in input_dir.glob(pattern) if p.is_file())
    if not config_files:
        print(f"No files matching '{pattern}' found in {input_dir.resolve()}")
        return
    for config_file in config_files:
        print(f"Processing {config_file}")
        try:
            process_config_file(config_file, output_dir, encoding, formats)
        except Exception as exc:  # keep processing remaining files
            print(f"Error processing {config_file}: {exc}")


def main(argv: list | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="huawei_policy_parse",
        description="Parse Huawei USG firewall configs "
                    "(address-sets, service-sets, security policies) to Excel.")
    parser.add_argument("-c", "--config", default=None,
                        help="Path to config.yaml "
                             "(default: config.yaml next to this tool).")
    args = parser.parse_args(argv)
    run(load_config(args.config))


if __name__ == "__main__":
    sys.exit(main())
