"""Parse Cisco ASA firewall configs: network/service object-groups and ACLs.

Refactored from FWConfigParse/cisco_ac_v7.py (tested against ASA 7.2(4) style
configs). Settings come from config.yaml instead of the missing config module
in the original project.
"""

from __future__ import annotations

import argparse
import datetime
import ipaddress
import json
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

from .service_port_mapping import SERVICE_TO_PORT

DEFAULT_CONFIG = {
    "input": {"directory": "Config", "pattern": "*.txt", "encoding": "utf-8"},
    "output": {"directory": "Output", "formats": ["excel", "json"]},
    "parse": {"replace_object_group": True},
}


def load_config(path: str | None = None) -> dict:
    config_path = Path(path) if path else Path(__file__).with_name("config.yaml")
    if not config_path.exists():
        return DEFAULT_CONFIG
    with open(config_path, encoding="utf-8") as fh:
        loaded = yaml.safe_load(fh) or {}
    return {section: {**DEFAULT_CONFIG[section], **loaded.get(section, {})}
            for section in DEFAULT_CONFIG}


def add_mask_to_ip(ip_str: str) -> str:
    """Add a host mask (/32 or /128) to a bare IP address string."""
    try:
        return f"{ipaddress.IPv4Address(ip_str)}/32"
    except ipaddress.AddressValueError:
        try:
            return f"{ipaddress.IPv6Address(ip_str)}/128"
        except ipaddress.AddressValueError:
            return ip_str


def parse_network_groups(lines: list) -> dict:
    """Parse ``object-group network`` stanzas into {group_name: [networks]}."""
    network_dict = {}
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if line.startswith("object-group network "):
            group_name = line.split("object-group network ")[-1].strip()
            network_dict[group_name] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("object-group") \
                    and not lines[i].startswith("!"):
                subline = lines[i].strip()
                if subline.startswith("network-object "):
                    parts = subline[len("network-object "):].split()
                    if len(parts) == 1:
                        network_dict[group_name].append(add_mask_to_ip(parts[0]))
                    elif parts[0] == 'host':
                        network_dict[group_name].append(add_mask_to_ip(parts[1]))
                    else:
                        try:
                            net = ipaddress.IPv4Network((parts[0], parts[1]),
                                                        strict=False)
                            network_dict[group_name].append(f"{net}")
                        except (ipaddress.AddressValueError,
                                ipaddress.NetmaskValueError):
                            try:
                                net = ipaddress.IPv6Network(f"{parts[0]}/{parts[1]}",
                                                            strict=False)
                                network_dict[group_name].append(f"{net}")
                            except ipaddress.AddressValueError:
                                network_dict[group_name].append(subline)
                i += 1
            continue

        i += 1

    return network_dict


def parse_service_groups(lines: list) -> dict:
    """Parse ``object-group service`` stanzas into
    {group_name: {"protocol": ..., "ports": [...]}}."""
    service_dict = {}
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if line.startswith("object-group service "):
            parts = line.split("object-group service ")[-1].strip().split()
            group_name = parts[0]
            protocol = parts[1] if len(parts) > 1 else ""
            service_dict[group_name] = {"protocol": protocol, "ports": []}
            i += 1
            while i < len(lines) and not lines[i].startswith("object-group") \
                    and not lines[i].startswith("!"):
                subline = lines[i].strip()
                if subline.startswith("port-object "):
                    port_parts = subline[len("port-object "):].split()
                    if port_parts[0] == 'eq':
                        service_name = port_parts[1]
                        port = SERVICE_TO_PORT.get(service_name, service_name)
                        service_dict[group_name]["ports"].append(port)
                    elif port_parts[0] == 'range':
                        service_dict[group_name]["ports"].append(
                            f"{port_parts[1]}-{port_parts[2]}")
                i += 1
            continue

        i += 1

    return service_dict


def parse_address(address_parts: list) -> tuple:
    """Parse a source/destination address; returns (address, words consumed)."""
    if not address_parts:
        return "", 0
    if address_parts[0] == 'host':
        return f"{address_parts[1]}/32", 2
    if address_parts[0] == 'object-group':
        return f"object-group: {address_parts[1]}", 2
    if address_parts[0] == 'any':
        return "any", 1
    if len(address_parts) > 1 and address_parts[1] in ('0.0.0.0', '255.255.255.255'):
        mask = address_parts[1]
        ip = address_parts[0]
        if mask == '0.0.0.0':
            return "any", 2
        return f"{ip}/32", 2
    ip = address_parts[0]
    mask = address_parts[1] if len(address_parts) > 1 else ""
    try:
        prefix = sum(bin(int(x)).count('1') for x in mask.split('.'))
        return f"{ip}/{prefix}", 2
    except (ValueError, AttributeError):
        return f"{ip} {mask}".strip(), 2 if mask else 1


def parse_port(port_parts: list) -> tuple:
    """Parse a destination port; returns (port, words consumed)."""
    if not port_parts:
        return "", 0
    if port_parts[0] == 'eq':
        if len(port_parts) > 1:
            port = SERVICE_TO_PORT.get(port_parts[1].lower(), port_parts[1])
            return str(port), 2
        return "", 1
    if port_parts[0] == 'range':
        if len(port_parts) > 2:
            port1 = SERVICE_TO_PORT.get(port_parts[1].lower(), port_parts[1])
            port2 = SERVICE_TO_PORT.get(port_parts[2].lower(), port_parts[2])
            return f"{port1}-{port2}", 3
        return "", 1
    if port_parts[0] == 'object-group':
        if len(port_parts) > 1:
            return f"object-group: {port_parts[1]}", 2
        return "", 1
    return "", 0


def parse_acl_entry(config_text):
    """Parse ACL entries from a single line, a list of lines, or full text."""
    if isinstance(config_text, list):
        config_text = "\n".join(config_text)
    if not isinstance(config_text, str):
        return None
    if '\n' in config_text:
        acl_rules = []
        for line in config_text.splitlines():
            line = line.strip()
            if line.startswith("access-list"):
                rule = parse_acl_entry(line)
                if rule:
                    acl_rules.append(rule)
        return acl_rules
    if not config_text.startswith("access-list"):
        return None
    parts = config_text.split()
    rule = {
        'name': parts[1],
        'type': parts[2],
        'action': parts[3],
        'protocol': parts[4],
    }
    src_index = 5
    rule['src_address'], src_words = parse_address(parts[src_index:])
    dst_index = src_index + src_words
    rule['dst_address'], dst_words = parse_address(parts[dst_index:])
    port_index = dst_index + dst_words
    rule['dst_port'], port_words = parse_port(parts[port_index:])
    other_index = port_index + port_words
    rule['other_config'] = ' '.join(parts[other_index:]) \
        if other_index < len(parts) else ""
    return rule


def expand_object_group_refs(acl_rules: list, networks: dict, services: dict) -> list:
    """Expand object-group references in ACL rules into newline-joined values."""
    expanded_rules = []
    for rule in acl_rules:
        new_rule = rule.copy()
        if rule['src_address'].startswith('object-group:'):
            group_name = rule['src_address'].split(':')[1].strip()
            if group_name in networks:
                new_rule['src_address'] = '\n'.join(networks[group_name])
        if rule['dst_address'].startswith('object-group:'):
            group_name = rule['dst_address'].split(':')[1].strip()
            if group_name in networks:
                new_rule['dst_address'] = '\n'.join(networks[group_name])
        if rule['dst_port'].startswith('object-group:'):
            group_name = rule['dst_port'].split(':')[1].strip()
            if group_name in services \
                    and services[group_name]['protocol'] == rule['protocol']:
                new_rule['dst_port'] = '\n'.join(services[group_name]['ports'])
        expanded_rules.append(new_rule)
    return expanded_rules


def write_to_excel(output_path: Path, networks: dict, services: dict,
                   acl_rules: list) -> None:
    """Write network groups, service groups and ACL rules to an Excel file."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    network_sheet = wb.create_sheet("地址组")
    network_sheet.append(["地址组名称", "IP地址"])
    for group_name, ip_list in networks.items():
        network_sheet.append([group_name, "\n".join(ip_list)])

    service_sheet = wb.create_sheet("服务组")
    service_sheet.append(["服务组名称", "协议", "端口号"])
    for group_name, service_info in services.items():
        service_sheet.append([group_name, service_info["protocol"],
                              "\n".join(service_info["ports"])])

    acl_sheet = wb.create_sheet("ACL规则")
    acl_sheet.append(["规则名称", "类型", "动作", "协议", "源地址",
                      "目的地址", "目的端口", "其他配置"])
    for rule in acl_rules:
        acl_sheet.append([
            rule.get('name', ''),
            rule.get('type', ''),
            rule.get('action', ''),
            rule.get('protocol', ''),
            rule.get('src_address', ''),
            rule.get('dst_address', ''),
            rule.get('dst_port', ''),
            rule.get('other_config', ''),
        ])

    bold_font = openpyxl.styles.Font(bold=True)
    wrap_alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for cell in ws[1]:
            cell.font = bold_font
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value and '\n' in str(cell.value):
                    line_count = str(cell.value).count('\n') + 1
                    ws.row_dimensions[cell.row].height = 14 * line_count
                cell.alignment = wrap_alignment
        if sheet == "ACL规则":
            column_widths = {'A': 20, 'B': 15, 'C': 15, 'D': 15,
                             'E': 30, 'F': 30, 'G': 20, 'H': 30}
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
        else:
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    if not cell.value:
                        continue
                    value = str(cell.value)
                    if '\n' in value:
                        max_length = max(max_length,
                                         max(len(line) for line in value.split('\n')))
                    else:
                        max_length = max(max_length, len(value))
                adjusted_width = (max_length + 2) * 1.2
                ws.column_dimensions[column_letter].width = min(adjusted_width, 50)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def export_to_json(output_path: Path, device_name: str, networks: dict,
                   services: dict, acl_rules: list) -> Path:
    """Export structured config data in the standard rule JSON format
    (consumed by Emulation/traffic_emulator; see Emulation/README.md)."""
    config_data = {
        "metadata": {
            "device_name": device_name,
            "vendor": "cisco",
            "export_time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        },
        "network_groups": networks,
        "service_groups": services,
        "rules": acl_rules,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(config_data, fh, ensure_ascii=False, indent=2)
    return output_path


def parse_config_text(config_content: str,
                      replace_object_group: bool = True) -> tuple:
    """Parse one config text; returns (networks, services, acl_rules)."""
    lines = config_content.splitlines()
    networks = parse_network_groups(lines)
    services = parse_service_groups(lines)
    acl_rules = parse_acl_entry(config_content) or []
    if replace_object_group:
        acl_rules = expand_object_group_refs(acl_rules, networks, services)
    return networks, services, acl_rules


def process_config_file(file_path: Path, output_dir: Path, formats: list,
                        encoding: str = "utf-8",
                        replace_object_group: bool = True) -> None:
    """Parse one Cisco config file and write Excel/JSON outputs."""
    config_content = file_path.read_text(encoding=encoding)
    networks, services, acl_rules = parse_config_text(
        config_content, replace_object_group)
    device_name = file_path.stem

    if "excel" in formats:
        excel_path = output_dir / f"{device_name}.xlsx"
        write_to_excel(excel_path, networks, services, acl_rules)
        print(f"Writing file {excel_path}")
    if "json" in formats:
        json_path = output_dir / f"{device_name}.json"
        export_to_json(json_path, device_name, networks, services, acl_rules)
        print(f"Writing file {json_path}")


def run(config: dict) -> None:
    input_dir = Path(config["input"]["directory"])
    pattern = config["input"].get("pattern", "*.txt")
    encoding = config["input"].get("encoding", "utf-8")
    output_dir = Path(config["output"]["directory"])
    formats = config["output"].get("formats", ["excel", "json"])
    replace_object_group = config["parse"].get("replace_object_group", True)

    if not input_dir.exists():
        print(f"Input directory not found: {input_dir.resolve()}")
        return
    config_files = sorted(input_dir.glob(pattern))
    if not config_files:
        print(f"No files matching '{pattern}' found in {input_dir.resolve()}")
        return
    for config_file in config_files:
        print(f"Processing {config_file}")
        process_config_file(config_file, output_dir, formats,
                            encoding, replace_object_group)


def main(argv: list | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="cisco_acl_parse",
        description="Parse Cisco ASA configs (object-groups and ACLs) "
                    "to Excel and JSON.")
    parser.add_argument("-c", "--config", default=None,
                        help="Path to config.yaml "
                             "(default: config.yaml next to this tool).")
    args = parser.parse_args(argv)
    run(load_config(args.config))


if __name__ == "__main__":
    sys.exit(main())
