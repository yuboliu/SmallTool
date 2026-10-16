import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from FW.cisco_acl_parse.converter import (
    expand_object_group_refs,
    load_config,
    parse_acl_entry,
    parse_config_text,
    parse_network_groups,
    parse_service_groups,
    process_config_file,
)

SAMPLE = Path(__file__).parent / "samples" / "asa-config.txt"


@pytest.fixture(scope="module")
def parsed():
    return parse_config_text(SAMPLE.read_text(encoding="utf-8"),
                             replace_object_group=False)


def test_network_groups(parsed):
    networks, _, _ = parsed
    assert set(networks) == {"OG-INSIDE", "OG-SERVERS"}
    assert networks["OG-INSIDE"] == ["10.1.0.0/24", "10.1.1.5/32"]
    assert networks["OG-SERVERS"] == ["172.16.0.0/16"]


def test_service_groups(parsed):
    _, services, _ = parsed
    assert services["OG-WEB"]["protocol"] == "tcp"
    assert services["OG-WEB"]["ports"] == ["80", "443", "8000-8010"]


def test_acl_entry_single_line():
    rule = parse_acl_entry(
        "access-list ACL extended permit tcp host 1.1.1.1 host 2.2.2.2 eq https")
    assert rule["name"] == "ACL"
    assert rule["type"] == "extended"
    assert rule["action"] == "permit"
    assert rule["protocol"] == "tcp"
    assert rule["src_address"] == "1.1.1.1/32"
    assert rule["dst_address"] == "2.2.2.2/32"
    assert rule["dst_port"] == "443"


def test_acl_rules_unexpanded(parsed):
    _, _, rules = parsed
    assert len(rules) == 3
    assert rules[0]["src_address"] == "any"
    assert rules[0]["dst_address"] == "object-group: OG-INSIDE"
    assert rules[0]["dst_port"] == "object-group: OG-WEB"
    assert rules[1]["src_address"] == "1.2.3.4/32"
    assert rules[1]["dst_address"] == "any"
    assert rules[2]["src_address"] == "192.168.0.0/24"
    assert rules[2]["dst_address"] == "203.0.113.5/32"
    assert rules[2]["dst_port"] == "22"


def test_expand_object_group_refs(parsed):
    networks, services, rules = parsed
    expanded = expand_object_group_refs(rules, networks, services)
    assert expanded[0]["dst_address"] == "10.1.0.0/24\n10.1.1.5/32"
    assert expanded[0]["dst_port"] == "80\n443\n8000-8010"
    # rules without object-group references stay unchanged
    assert expanded[1] == rules[1]


def test_process_config_file_writes_excel_and_json(tmp_path):
    process_config_file(SAMPLE, tmp_path, ["excel", "json"])

    excel_path = tmp_path / "asa-config.xlsx"
    json_path = tmp_path / "asa-config.json"
    assert excel_path.exists()
    assert json_path.exists()

    workbook = load_workbook(excel_path)
    assert workbook.sheetnames == ["地址组", "服务组", "ACL规则"]
    acl_sheet = workbook["ACL规则"]
    assert acl_sheet.cell(row=1, column=1).value == "规则名称"
    assert acl_sheet.cell(row=2, column=1).value == "OUTSIDE-IN"
    # expanded destination address of the first rule
    assert "10.1.0.0/24" in acl_sheet.cell(row=2, column=6).value

    with open(json_path, encoding="utf-8") as fh:
        data = json.load(fh)
    assert data["metadata"]["device_name"] == "asa-config"
    assert data["metadata"]["vendor"] == "cisco"
    assert len(data["rules"]) == 3
    assert data["rules"][0]["dst_address"] == "10.1.0.0/24\n10.1.1.5/32"
    assert data["network_groups"]["OG-INSIDE"] == ["10.1.0.0/24", "10.1.1.5/32"]


def test_run_with_yaml_config(tmp_path, monkeypatch):
    config_dir = tmp_path / "Config"
    config_dir.mkdir()
    (config_dir / "asa.txt").write_text(SAMPLE.read_text(encoding="utf-8"),
                                        encoding="utf-8")
    config_file = tmp_path / "config.yaml"
    config_file.write_text(
        "input:\n  directory: Config\n"
        "output:\n  directory: Out\n  formats: [excel]\n",
        encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    from FW.cisco_acl_parse.converter import run
    run(load_config(str(config_file)))
    assert (tmp_path / "Out" / "asa.xlsx").exists()
    assert not (tmp_path / "Out" / "asa.json").exists()
