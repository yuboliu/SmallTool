import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from FW.huawei_policy_parse.converter import (
    convert_text,
    load_config,
    process_config_file,
    run,
)
from FW.huawei_policy_parse.extractors import (
    extract_address_sets,
    extract_security_policies,
    extract_service_sets,
)

SAMPLE = Path(__file__).parent / "samples" / "usg-config.txt"


@pytest.fixture(scope="module")
def content():
    return SAMPLE.read_text(encoding="utf-8")


def test_address_sets(content):
    address_sets = extract_address_sets(content)
    assert len(address_sets) == 2
    inside = address_sets[0]
    assert inside["set_name"] == "INSIDE-NET"
    assert inside["set_type"] == "object"
    assert inside["addresses"] == (
        "10.1.0.0/24\n192.168.1.10/32\n172.16.1.1 - 172.16.1.10")


def test_service_sets(content):
    service_sets = extract_service_sets(content.splitlines())
    assert len(service_sets) == 1
    web = service_sets[0]
    assert web["service_set_name"] == "WEB-SVC"
    assert web["services"] == "tcp 80\ntcp 443\ntcp 8000 to 8010"


def test_security_policies(content):
    policies = extract_security_policies(content.splitlines())
    assert len(policies) == 2
    allow = policies[0]
    assert allow["rule_name"] == "ALLOW-WEB"
    assert allow["source_zone"] == "trust"
    assert allow["destination_zone"] == "untrust"
    assert allow["source_address"] == "INSIDE-NET"
    assert allow["destination_address"] == "203.0.113.10 mask 32"
    assert allow["service"] == "WEB-SVC"
    assert allow["action"] == "permit"
    deny = policies[1]
    assert deny["source_address"] == "DMZ-NET"
    assert deny["destination_address"] == "INSIDE-NET"
    assert deny["action"] == "deny"


def test_convert_text_sheets(content):
    sheets = convert_text(content)
    assert list(sheets) == ["地址集", "服务集", "安全策略-原始", "安全策略-关联"]
    _, correlated = sheets["安全策略-关联"]
    allow = correlated[0]
    assert allow["source_address"] == (
        "10.1.0.0/24\n192.168.1.10/32\n172.16.1.1 - 172.16.1.10")
    # non-set destination stays unchanged
    assert allow["destination_address"] == "203.0.113.10 mask 32"
    assert allow["service"] == "tcp 80\ntcp 443\ntcp 8000 to 8010"
    deny = correlated[1]
    assert deny["source_address"] == "172.18.0.0/24"


def test_process_config_file_writes_excel(tmp_path):
    output_path = process_config_file(SAMPLE, tmp_path)
    assert output_path.exists()
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["地址集", "服务集", "安全策略-原始", "安全策略-关联"]
    sheet = workbook["安全策略-关联"]
    assert sheet.cell(row=1, column=1).value == "规则名称"
    assert sheet.cell(row=2, column=1).value == "ALLOW-WEB"


def test_process_config_file_writes_standard_json(tmp_path):
    process_config_file(SAMPLE, tmp_path, formats=["json"])
    json_path = tmp_path / "usg-config.json"
    assert json_path.exists()
    with open(json_path, encoding="utf-8") as fh:
        data = json.load(fh)
    assert data["metadata"]["device_name"] == "usg-config"
    assert data["metadata"]["vendor"] == "huawei"
    assert len(data["address_sets"]) == 2

    rules = data["rules"]
    assert len(rules) == 2
    allow = rules[0]
    assert allow["name"] == "ALLOW-WEB"
    assert allow["action"] == "permit"
    assert allow["protocol"] == "tcp"
    assert allow["dst_port"] == "80\n443\n8000-8010"
    assert allow["src_address"] == (
        "10.1.0.0/24\n192.168.1.10/32\n172.16.1.1 - 172.16.1.10")
    assert allow["dst_address"] == "203.0.113.10 mask 32"
    assert allow["source_zone"] == "trust"
    deny = rules[1]
    assert deny["protocol"] == "tcp"
    assert deny["dst_port"] == "3389"
    assert deny["src_address"] == "172.18.0.0/24"


def test_run_with_yaml_config(tmp_path, monkeypatch):
    config_dir = tmp_path / "Config"
    config_dir.mkdir()
    (config_dir / "usg.txt").write_text(SAMPLE.read_text(encoding="utf-8"),
                                        encoding="utf-8")
    config_file = tmp_path / "config.yaml"
    config_file.write_text(
        "input:\n  directory: Config\noutput:\n  directory: Out\n",
        encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    run(load_config(str(config_file)))
    assert (tmp_path / "Out" / "usg.xlsx").exists()
