from pathlib import Path

import pytest
from openpyxl import load_workbook

from LB.adp_conf_to_excel.converter import (
    SHEET_ORDER,
    build_sheets,
    decode_unicode,
    load_config,
    run,
)

SAMPLE = Path(__file__).parent / "samples" / "startup-config"


@pytest.fixture(scope="module")
def sheets():
    return build_sheets(SAMPLE.read_text(encoding="utf-8"))


def rows(sheets, sheet):
    return sheets[sheet][1]


def test_decode_unicode():
    assert decode_unicode("%u6D4B%u8BD5") == "测试"
    assert decode_unicode("abc%u6D4Bdef") == "abc测def"
    assert decode_unicode("plain") == "plain"


def test_ip_nat_pool(sheets):
    (nat,) = rows(sheets, "ip_nat_pool")
    assert nat["Name"] == "natpool1"
    assert nat["start"] == "10.1.1.1"
    assert nat["end"] == "10.1.1.10"
    assert nat["netmask"] == "255.255.255.0"


def test_health_check(sheets):
    (hc,) = rows(sheets, "health_check")
    assert hc["Name"] == "hc_tcp"
    assert hc["interval"] == "5"
    assert hc["retry"] == "3"
    assert hc["timeout"] == "3"
    assert hc["up-check-cnt"] == "2"
    assert hc["method"] == "tcp"
    assert hc["wait-all-retry"] is True


def test_slb_pool(sheets):
    (pool,) = rows(sheets, "slb_pool")
    assert pool["Name"] == "pool_web"
    assert pool["protocol"] == "tcp"
    assert pool["member"] == "192.168.1.1:80\r\n192.168.1.2:80"
    assert pool["health-check"] == "hc_tcp"


def test_slb_profile(sheets):
    (profile,) = rows(sheets, "slb_profile")
    assert profile["Name"] == "tcp_prof"
    assert profile["type"] == "tcp"
    assert profile["idle-timeout"] == "300"
    assert profile["tcp-keepalive"] is True


def test_slb_profile_persist(sheets):
    (persist,) = rows(sheets, "slb_profile_persist")
    assert persist["Name"] == "persist_src"
    assert persist["timeout"] == "600"


def test_slb_virtual_address(sheets):
    vas = rows(sheets, "slb_virtual-address")
    assert len(vas) == 2
    assert vas[0]["Virtual_Address"] == "10.10.10.10"
    assert vas[0]["Port"] == "port 80 tcp"
    assert vas[0]["pool"] == "pool_web"
    assert vas[0]["persist"] == "persist_src"
    assert vas[0]["description"] == "测试"
    assert vas[0]["Other"] == "enable"
    assert vas[1]["Port"] == "port 443 tcp"
    assert "persist" not in vas[1]


def test_run_writes_excel(tmp_path, monkeypatch):
    input_file = tmp_path / "startup-config"
    input_file.write_text(SAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    config_file = tmp_path / "config.yaml"
    config_file.write_text(
        'input:\n  file: "startup-config"\noutput:\n  file: "out.xlsx"\n',
        encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    output_path = run(load_config(str(config_file)))
    assert output_path.exists()
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == SHEET_ORDER
    sheet = workbook["slb_virtual-address"]
    assert sheet.cell(row=1, column=1).value == "Virtual_Address"
    assert sheet.cell(row=2, column=1).value == "10.10.10.10"
