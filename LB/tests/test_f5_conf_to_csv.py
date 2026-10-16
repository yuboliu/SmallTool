import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from LB.f5_conf_to_csv.converter import (
    SHEET_ORDER,
    convert_file,
    extract_tables,
    load_config,
    run,
)

SAMPLE = Path(__file__).parent / "samples" / "bigip.conf"


@pytest.fixture(scope="module")
def tables():
    return extract_tables(SAMPLE.read_text(encoding="utf-8"))


def rows(tables, sheet):
    return tables[sheet][1]


def test_nodes(tables):
    nodes = rows(tables, "f5-nodes")
    assert {n["Name"] for n in nodes} == {"/Common/node_web1", "/Common/node_web2"}
    assert nodes[0]["Address"] == "10.0.0.1"


def test_pool_members(tables):
    pools = rows(tables, "f5-pools")
    web = [p for p in pools if p["Name"] == "/Common/pool_web"]
    assert len(web) == 2
    assert web[0]["Member_Name"] == "/Common/node_web1:80"
    assert web[0]["Member_Address"] == "10.0.0.1"
    assert web[0]["Monitor"] == "/Common/mon_http"
    assert web[0]["LB_Mode"] == "round-robin"
    empty = [p for p in pools if p["Name"] == "/Common/pool_empty"]
    assert len(empty) == 1
    assert empty[0]["Member_Name"] == ""


def test_virtual_server(tables):
    (vs,) = rows(tables, "f5-virtualservers")
    assert vs["Name"] == "/Common/vs_web"
    assert vs["Destination"] == "/Common/va_web:80"
    assert vs["IP_Protocol"] == "tcp"
    assert vs["Pool"] == "/Common/pool_web"
    assert vs["Profiles"] == "/Common/http; /Common/tcp"
    assert vs["Rules"] == "/Common/rule_redirect"
    assert vs["Persist"] == "/Common/persist_src"
    assert "automap" in vs["Source_Address_Translation"]
    assert vs["Translate_Address"] == "enabled"


def test_virtual_address(tables):
    (va,) = rows(tables, "f5-virtualaddresses")
    assert va["Address"] == "192.168.1.10"
    assert va["Arp"] == "enabled"
    assert va["Traffic_Group"] == "/Common/traffic-group-1"


def test_monitor(tables):
    (mon,) = rows(tables, "f5-monitors")
    assert mon["Name"] == "/Common/mon_http"
    assert mon["Type"] == "http"
    assert mon["Interval"] == "5"
    assert mon["Timeout"] == "16"
    assert mon["Send"] == '"GET / HTTP/1.0\\r\\n\\r\\n"'


def test_profile(tables):
    (prof,) = rows(tables, "f5-profiles")
    assert prof["Name"] == "/Common/ssl_web"
    assert prof["Type"] == "client-ssl"
    assert prof["Cert"] == "/Common/web.crt"
    assert prof["Key"] == "/Common/web.key"


def test_persistence(tables):
    (persist,) = rows(tables, "f5-persistence")
    assert persist["Name"] == "/Common/persist_src"
    assert persist["Type"] == "source-addr"
    assert persist["Timeout"] == "180"


def test_rule_raw_text(tables):
    (rule,) = rows(tables, "f5-rules")
    assert rule["Name"] == "/Common/rule_redirect"
    assert "when HTTP_REQUEST" in rule["Rule"]
    assert "HTTP::redirect" in rule["Rule"]


def test_overview(tables):
    overview = rows(tables, "f5-overview")
    web = [o for o in overview if o["Pool"] == "/Common/pool_web"]
    assert len(overview) == 3  # 2 members joined to vs + 1 empty pool
    assert all(o["VirtualServer"] == "/Common/vs_web" for o in web)
    assert {o["Pool Member"] for o in web} == {
        "/Common/node_web1:80", "/Common/node_web2:80"}
    empty = [o for o in overview if o["Pool"] == "/Common/pool_empty"]
    assert empty[0]["VirtualServer"] == ""


def test_convert_file_writes_csv_and_excel(tmp_path):
    out_dir = convert_file(SAMPLE, tmp_path, ["csv", "excel"])
    for sheet in SHEET_ORDER:
        csv_path = out_dir / f"{sheet}.csv"
        assert csv_path.exists()
        with open(csv_path, encoding="utf-8-sig") as fh:
            assert list(csv.DictReader(fh)) is not None
    excel_path = out_dir / "bigip.xlsx"
    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    assert workbook.sheetnames == SHEET_ORDER
    sheet = workbook["f5-pools"]
    assert sheet.cell(row=1, column=1).value == "Name"
    assert sheet.cell(row=2, column=1).value == "/Common/pool_web"


def test_run_with_yaml_config(tmp_path, monkeypatch):
    (tmp_path / "bigip.conf").write_text(SAMPLE.read_text(encoding="utf-8"),
                                         encoding="utf-8")
    config_file = tmp_path / "config.yaml"
    config_file.write_text(
        "input:\n  directory: .\noutput:\n  directory: Output\n",
        encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    run(load_config(str(config_file)))
    assert (tmp_path / "Output" / "bigip" / "f5-nodes.csv").exists()
    assert (tmp_path / "Output" / "bigip" / "bigip.xlsx").exists()
